"""Populate the team's AS9102 template (Forms 1-3) - values only, template formatting preserved."""
import io
import json
from copy import copy
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = ROOT / "templates" / "AS9102_template.xlsx"

# Cell map for templates/AS9102_template.xlsx - edit here if the template layout changes.
FORM1 = {"sheet": 0, "cells": {"part_number": "B3", "part_name": "D3", "serial_number": "F3",
                               "fai_report_number": "B4", "revision": "D4", "drawing_number": "F4",
                               "drawing_revision": "B5", "prepared_by": "B10", "date": "D10"}}
FORM2 = {"sheet": 1, "row": 4, "material": "A", "spec": "B"}
FORM3 = {"sheet": 2, "start_row": 4, "cols": {"char_no": "A", "location": "B", "designator": "C",
                                               "requirement": "D", "results": "E", "tooling": "F",
                                               "nonconformance": "G", "status": "H"}}
HEADER_FILL = PatternFill("solid", fgColor="203764")


def _set(ws, ref, value):
    for rng in ws.merged_cells.ranges:
        if ref in rng:
            ref = rng.start_cell.coordinate
            break
    ws[ref] = value


def _copy_row_style(ws, src_row: int, dst_row: int, max_col: int):
    for col in range(1, max_col + 1):
        s, d = ws.cell(src_row, col), ws.cell(dst_row, col)
        if s.has_style:
            d.font, d.border, d.fill = copy(s.font), copy(s.border), copy(s.fill)
            d.alignment, d.number_format = copy(s.alignment), s.number_format
    if ws.row_dimensions[src_row].height:
        ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def requirement(row: dict) -> str:
    unit = row.get("unit") or ""
    dim_type = row.get("type", "").lower() in ("linear", "diameter", "radius", "chamfer", "basic")
    tail = f" {unit}" if dim_type and unit and unit not in row["notation"] and unit.lower() not in ("deg", "°") else ""
    qty = f" ({row['qty']}X)" if (row.get("qty") or 1) > 1 and f"{row['qty']}X" not in row["notation"].upper() else ""
    return f"{row['name']}: {row['notation']}{tail}{qty}"


def _num(r):
    try:
        return float(r["id"])
    except ValueError:
        return 1e9


def exported_rows(rows):
    return sorted((r for r in rows if r["status"] in ("Accepted", "Corrected")), key=_num)


def build_workbook(rows: list[dict], drawing, info: dict, balloon_png: bytes | None,
                   template: bytes | None = None, model: str = "") -> bytes:
    wb = load_workbook(io.BytesIO(template) if template else DEFAULT_TEMPLATE)
    sheets = wb.worksheets

    # Form 1 - part accountability
    if len(sheets) >= 3:
        f1 = sheets[FORM1["sheet"]]
        for key, ref in FORM1["cells"].items():
            if info.get(key):
                _set(f1, ref, info[key])
        # Form 2 - material
        f2 = sheets[FORM2["sheet"]]
        if info.get("material"):
            _set(f2, f"{FORM2['material']}{FORM2['row']}", info["material"])
            _set(f2, f"{FORM2['spec']}{FORM2['row']}", info.get("material_spec") or "")
        f3 = sheets[FORM3["sheet"]]
    else:
        f3 = sheets[0]

    # Form 3 - characteristics
    cols, start = FORM3["cols"], FORM3["start_row"]
    sheet_no = drawing.page + 1
    for i, r in enumerate(exported_rows(rows)):
        n = start + i
        if i:
            _copy_row_style(f3, start, n, 8)
        _set(f3, f"{cols['char_no']}{n}", r["id"])
        _set(f3, f"{cols['location']}{n}", f"Sht {sheet_no}, {r['zone']}" if r.get("zone") else f"Sht {sheet_no}")
        _set(f3, f"{cols['designator']}{n}", r.get("designator") or "")
        _set(f3, f"{cols['requirement']}{n}", requirement(r))
        _set(f3, f"{cols['tooling']}{n}", r.get("tool") or "")

    # Traceability sheet
    ts = wb.create_sheet("Traceability")
    thin = Side(style="thin", color="BFBFBF")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    meta = [("Drawing file", drawing.filename), ("Drawing SHA-256", drawing.sha256), ("Sheet / page", sheet_no),
            ("Image size (px)", f"{drawing.image.width} x {drawing.image.height}"), ("AI model", model),
            ("Exported (UTC)", now),
            ("Statement", "AI-assisted extraction. Every exported characteristic was accepted or corrected by a reviewer.")]
    for i, (k, v) in enumerate(meta, 1):
        ts.cell(i, 1, k).font = Font(bold=True)
        ts.cell(i, 2, str(v))
    hdr = len(meta) + 2
    heads = ["Char No.", "Balloon on drawing", "Zone", "Review status", "In Form 3", "Requirement (final)",
             "Notation (AI)", "Feature (AI)", "Confidence", "Flags", "Edited fields", "Location px (x0,y0,x1,y1)",
             "Location exact?", "Balloon px (cx,cy,r)", "Text tokens", "AI raw"]
    for j, h in enumerate(heads, 1):
        c = ts.cell(hdr, j, h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    exp = {r["key"] for r in exported_rows(rows)}
    for i, r in enumerate(sorted(rows, key=_num), hdr + 1):
        ai = r.get("_ai") or {}
        vals = [r["id"], r["balloon"] or "generated", r.get("zone", ""), r["status"], "yes" if r["key"] in exp else "no",
                requirement(r), ai.get("notation", ""), ai.get("name", ""), r.get("confidence"),
                "; ".join(r["flags"]), ", ".join(r.get("_edited", [])), str(r.get("loc") or ""),
                "yes" if r.get("exact") else "approx", str(r.get("balloon_xy") or ""), " ".join(r.get("tokens", [])),
                json.dumps(ai, ensure_ascii=False)]
        for j, v in enumerate(vals, 1):
            c = ts.cell(i, j, v)
            c.border = Border(bottom=thin)
            c.alignment = Alignment(wrap_text=j in (6, 10), vertical="top")
    for col, w in zip("ABCDEFGHIJKLMNOP", [9, 10, 7, 12, 8, 44, 24, 24, 10, 50, 16, 22, 10, 16, 14, 40]):
        ts.column_dimensions[col].width = w
    ts.freeze_panes = ts.cell(hdr + 1, 1)

    if balloon_png:
        ds = wb.create_sheet("Ballooned Drawing")
        img = XLImage(io.BytesIO(balloon_png))
        ratio = 1500 / img.width
        img.width, img.height = int(img.width * ratio), int(img.height * ratio)
        ds.add_image(img, "A1")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
