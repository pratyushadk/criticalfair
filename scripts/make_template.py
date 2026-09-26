"""Build templates/AS9102_template.xlsx from the team's filled example: keep labels + formatting, clear data."""
import sys
from pathlib import Path
from openpyxl import load_workbook

src = Path(sys.argv[1] if len(sys.argv) > 1 else "~/Desktop/AS9102_FAIR_Cylinder.xlsx").expanduser()
wb = load_workbook(src)
f1, f2, f3 = wb.worksheets[:3]
for ref in ["B3", "D3", "F3", "B4", "D4", "F4", "B5", "D5", "F5", "B6", "D6", "F6", "B7", "D7", "F7", "B8", "B9",
            "B10", "D10", "B11", "D11", "B12", "D12"]:
    f1[ref].value = None
for row in f2.iter_rows(min_row=4, max_row=f2.max_row):
    for c in row:
        c.value = None
for row in f3.iter_rows(min_row=4, max_row=f3.max_row):
    for c in row:
        c.value = None
out = Path(__file__).resolve().parent.parent / "templates" / "AS9102_template.xlsx"
wb.save(out)
print("wrote", out, [ws.title for ws in wb.worksheets])
