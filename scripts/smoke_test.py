"""Headless end-to-end UI test on a sample with a saved analysis (no API call)."""
import io
from pathlib import Path

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parent.parent
at = AppTest.from_file(str(ROOT / "app" / "main.py"), default_timeout=120).run()


def click(label, exact=False):
    next(b for b in at.button if (b.label == label if exact else b.label.startswith(label))).click().run()
    assert not at.exception, at.exception


for sample in ["flange.jpeg", "cylinder_ballooned.png"]:
    at.session_state["source"] = (sample, (ROOT / "samples" / "drawings" / sample).read_bytes())
    at.run()
    click("Analyse drawing")          # -> analyse stage runs to completion -> review
    assert at.session_state["step"] == 2, at.session_state["step"]
    rows = at.session_state["rows"]
    click("Accept all unflagged")
    while any(r["status"] == "Pending" for r in at.session_state["rows"]):
        click("Accept", exact=True)
    click("Continue to export")
    assert at.session_state["step"] == 3
    wb = load_workbook(io.BytesIO(at.session_state["xlsx"]))
    f3 = wb.worksheets[2]
    data = [r for r in f3.iter_rows(min_row=4, max_col=8, values_only=True) if r[0]]
    print(sample, "->", wb.sheetnames, f"{len(data)} Form 3 rows")
    for r in data[:4]:
        print("   ", r)
    print("    Form1:", wb.worksheets[0]["B3"].value, wb.worksheets[0]["D3"].value)
    click("New drawing")
print("OK")
