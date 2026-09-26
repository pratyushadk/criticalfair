"""Cached analysis -> accept all -> fill the AS9102 template; print what landed where."""
import io, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from openpyxl import load_workbook
from app import extractor, validation, overlay, as9102
from app.drawing import load_drawing
from app.perception import perceive
for p in sys.argv[1:]:
    d = load_drawing(Path(p).name, Path(p).read_bytes())
    P = perceive(d.image, d.pdf_text)
    ext = extractor.cached(d.sha256)
    rows, _ = validation.build(ext, P)
    overlay.place_balloons(d.image, rows, P)
    for r in rows: r["status"] = "Accepted"
    buf = io.BytesIO(); overlay.draw(d.image, rows, clean=True).save(buf, "PNG")
    x = as9102.build_workbook(rows, d, ext.drawing.model_dump() | {"drawing_revision": ext.drawing.revision}, buf.getvalue())
    out = Path("output") / f"{Path(p).stem}_AS9102.xlsx"; out.write_bytes(x)
    wb = load_workbook(out)
    print("==", out, wb.sheetnames)
    pi = wb["Part Info, Drawing & GDT Ref"]
    print("  part:", pi["B4"].value, "|", pi["D4"].value, "| rev", pi["B5"].value, "| mat", pi["B11"].value, "| equip", pi["D14"].value, "| img", len(pi._images))
    ch = wb["Characteristics"]
    for row in ch.iter_rows(min_row=1, max_row=6, max_col=11, values_only=True): print("  ", row)
    print("  dv:", [(str(v.sqref), v.formula1) for v in ch.data_validations.dataValidation])
    print("  summary:", wb["Summary"]["B3"].value, "| cf:", [str(c.sqref) for c in ch.conditional_formatting])
